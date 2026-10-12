from __future__ import annotations

import io
import re
from dataclasses import dataclass


@dataclass
class ResumeSnapshot:
    text: str
    school: str = ""
    major: str = ""
    degree: str = ""
    graduation_year: int | None = None
    skills: list[str] | None = None


def parse_resume(uploaded_file) -> ResumeSnapshot:
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()
    if name.endswith(".docx"):
        from docx import Document

        document = Document(io.BytesIO(raw))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.replace("\n", " ") for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith(".txt"):
        text = raw.decode("utf-8", errors="replace")
    else:
        raise ValueError("仅支持 DOCX、PDF 和 TXT 简历。")
    return snapshot(text)


_MAJOR_PATTERNS = [
    # 表格简历常见写法：学校 | 专业 | 学历
    r"(?:大学|学院)[^\n]{0,30}[|｜]\s*([^|｜\n]{2,18})\s*[|｜]",
    # 纯文本 / PDF 简历常见写法：「专业：经济统计学」或「专业 经济统计学」。
    # 旧实现只认表格分隔符，没有 | 的简历一律提取不到专业。
    r"^\s*(?:所学专业|专业|主修)\s*[：:]\s*([^\s|｜，,。；;\n]{2,18})",
    r"^\s*(?:所学专业|专业|主修)\s+([^\s|｜，,。；;\n]{2,18})",
]


def snapshot(text: str) -> ResumeSnapshot:
    from radar_core import extract_skills

    school_match = re.search(r"([\u4e00-\u9fff]{2,20}(?:大学|学院))", text)
    major_match = next(
        (match for pattern in _MAJOR_PATTERNS if (match := re.search(pattern, text, re.MULTILINE))),
        None,
    )
    degree_match = re.search(r"本科|硕士|博士|大专", text)
    # 「2028届」这种写法后面没有分隔符，旧正则 (20\d{2})[.年/-] 会整条漏掉，
    # 于是毕业年份回填不上。这里把「届」也算作合法的后接字符。
    years = [int(value) for value in re.findall(r"(20\d{2})\s*(?:届|[.年/-])", text)]
    graduation_year = max(years) if years else None
    return ResumeSnapshot(
        text=text.strip(),
        school=school_match.group(1) if school_match else "",
        major=major_match.group(1).strip() if major_match else "",
        degree=degree_match.group(0) if degree_match else "",
        graduation_year=graduation_year,
        skills=extract_skills(text),
    )
