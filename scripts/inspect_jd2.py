from pathlib import Path
from docx import Document

workspace = Path(__file__).resolve().parents[2]
source = workspace / "JD2.docx"
doc = Document(source)

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text:
        print(f"{index:04d}\t{text}")

for table_index, table in enumerate(doc.tables):
    print(f"\n[TABLE {table_index}] rows={len(table.rows)} cols={len(table.columns)}")
    for row in table.rows:
        print("\t".join(cell.text.replace("\n", " | ").strip() for cell in row.cells))
